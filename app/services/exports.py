"""Выгрузки по вкладкам: «Статус бота» (все отобранные со ссылками и статусами) и по вкладке
на каждый день чекина. Одни и те же таблицы пишутся в Google Sheet и в Excel-файл.
"""
import io
import re
from dataclasses import dataclass
from datetime import datetime
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Checkin, CheckinMethod, EventDay, Participant, Poll, PollAnswer, PollKind, PollOption
from app.services.importer import team_display_name
from app.services.tokens import invite_link

YES, NO = "да", "нет"
METHOD_TITLES = {CheckinMethod.SELF_QR: "QR", CheckinMethod.VOLUNTEER: "волонтёр", CheckinMethod.ADMIN: "админ"}


@dataclass
class Table:
    name: str
    header: list[str]
    rows: list[list[str]]


def _fmt(dt: datetime | None, tz: ZoneInfo, pattern: str = "%d.%m %H:%M") -> str:
    return dt.astimezone(tz).strftime(pattern) if dt else ""


def checkin_tab_name(day: EventDay) -> str:
    # Google Sheets не любит в названиях листов []:*?/\ и длину > 100.
    return re.sub(r"[\[\]:*?/\\]", " ", f"Чекин {day.label}")[:100]


async def _attendance(session: AsyncSession) -> dict[int, tuple[str, str | None, str]]:
    poll = await session.scalar(
        select(Poll).where(Poll.kind == PollKind.ATTENDANCE).order_by(Poll.created_at.desc()).limit(1)
    )
    if poll is None:
        return {}
    rows = await session.execute(
        select(PollAnswer.participant_id, PollOption.label_ru, PollAnswer.reason, PollAnswer.source)
        .join(PollOption, PollOption.id == PollAnswer.option_id)
        .where(PollAnswer.poll_id == poll.id)
    )
    return {pid: (label, reason, source) for pid, label, reason, source in rows}


async def passed_table(session: AsyncSession, bot_username: str, timezone: str, name: str = "Статус бота") -> Table:
    """Вкладка «Статус бота»: строка = отобранный участник. Отсюда берут ссылку для письма."""
    tz = ZoneInfo(timezone)
    days = (await session.scalars(select(EventDay).order_by(EventDay.date))).all()
    answers = await _attendance(session)
    checkins = {(c.participant_id, c.event_day_id): c.created_at for c in (await session.scalars(select(Checkin))).all()}
    participants = (
        await session.scalars(select(Participant).where(Participant.is_selected).order_by(Participant.id))
    ).all()

    header = [
        "ID заявки", "Имя", "Почта", "Телефон", "Telegram (форма)", "Команда", "Роль", "Вуз",
        "Ссылка-приглашение", "Открыл бота", "Telegram в боте", "Заблокировал бота", "Придёт?", "Причина",
        *[f"Чекин: {d.label}" for d in days],
    ]
    rows = []
    for p in participants:
        answer = answers.get(p.id)
        answer_text = ""
        if answer:
            answer_text = answer[0] + (" (обзвон)" if answer[2] == "call" else "")
        rows.append(
            [
                p.application_id or "",
                p.full_name,
                p.email,
                p.phone or p.phone_raw or "",
                f"@{p.tg_username_form}" if p.tg_username_form else "",
                team_display_name(p.team),
                p.team_role or "",
                p.university or "",
                invite_link(bot_username, p.invite_token) if bot_username else p.invite_token,
                YES if p.telegram_user_id else NO,
                f"@{p.telegram_username}" if p.telegram_username else ("" if not p.telegram_user_id else str(p.telegram_user_id)),
                YES if p.is_blocked_bot else "",
                answer_text,
                (answer[1] or "") if answer else "",
                *[_fmt(checkins.get((p.id, d.id)), tz, "%H:%M") for d in days],
            ]
        )
    return Table(name, header, rows)


async def checkin_table(session: AsyncSession, day: EventDay, timezone: str) -> Table:
    """Вкладка дня: кто отметился и когда, по времени."""
    tz = ZoneInfo(timezone)
    result = await session.execute(
        select(Checkin, Participant)
        .join(Participant, Participant.id == Checkin.participant_id)
        .where(Checkin.event_day_id == day.id)
        .order_by(Checkin.created_at)
    )
    header = ["Время", "Имя", "Команда", "Телефон", "Почта", "Telegram", "Способ", "Кто отметил (ID)"]
    rows = [
        [
            _fmt(c.created_at, tz, "%H:%M:%S"),
            p.full_name,
            team_display_name(p.team),
            p.phone or p.phone_raw or "",
            p.email,
            f"@{p.telegram_username}" if p.telegram_username else "",
            METHOD_TITLES.get(c.method, c.method.value),
            str(c.checked_by) if c.checked_by else "",
        ]
        for c, p in result.all()
    ]
    return Table(checkin_tab_name(day), header, rows)


async def all_tables(session: AsyncSession, bot_username: str, timezone: str, passed_name: str = "Статус бота") -> list[Table]:
    tables = [await passed_table(session, bot_username, timezone, passed_name)]
    for day in (await session.scalars(select(EventDay).order_by(EventDay.date))).all():
        tables.append(await checkin_table(session, day, timezone))
    return tables


async def bot_status_table(session: AsyncSession, bot_username: str, timezone: str) -> tuple[list[str], list[list[str]]]:
    """Вкладка «Статус бота» как (заголовок, строки) — для CSV."""
    table = await passed_table(session, bot_username, timezone)
    return table.header, table.rows


def tables_to_xlsx(tables: list[Table]) -> bytes:
    """Excel-файл: по листу на таблицу, жирная шапка, закреплённая первая строка, ширина по содержимому."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="5B42DC")
    for table in tables:
        ws = wb.create_sheet(table.name[:31])  # Excel ограничивает имя листа 31 символом
        ws.append(table.header)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for row in table.rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max((len(str(c.value or "")) for c in column), default=8)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

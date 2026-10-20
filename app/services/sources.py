"""Откуда берутся строки заявок: CSV (загрузка в бот, CLI) или Google Sheets.

Все обращения к Sheets идут в отдельном потоке (gspread синхронный) с таймаутом
и ретраями с экспоненциальной задержкой. Ошибка Sheets не роняет бота: вызывающий
код получает SheetsError и сообщает о ней админу.
"""
import asyncio
import base64
import csv
import io
import json
from collections.abc import Callable
from typing import Any, TypeVar

from tenacity import AsyncRetrying, retry_if_not_exception_type, stop_after_attempt, wait_exponential

from app.config import Settings
from app.logs import get_logger

log = get_logger(__name__)
T = TypeVar("T")
Row = dict[str, str]


class SheetsError(Exception):
    pass


class SheetMissing(SheetsError):
    """Нужной вкладки нет — повторять запрос бессмысленно."""


# ---------------------------------------------------------------- CSV

def parse_csv(data: bytes) -> list[Row]:
    """CSV из Excel/Google Sheets: UTF-8 с BOM или без, разделитель «,» или «;»."""
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не удалось прочитать файл: сохраните CSV в кодировке UTF-8")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for row in reader:
        clean = {(k or "").strip(): (v or "").strip() for k, v in row.items() if k is not None}
        if any(clean.values()):
            rows.append(clean)
    return rows


def parse_xlsx(data: bytes, sheet: str | None = None) -> list[Row]:
    """Excel-файл (.xlsx): берётся лист с нужным именем или первый. Первая строка — заголовок."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Не удалось открыть Excel-файл: {exc}") from exc
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    values = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    if not values:
        return []
    header = values[0]
    rows = []
    for line in values[1:]:
        row = {h: v for h, v in zip(header, line) if h}
        if any(row.values()):
            rows.append(row)
    return rows


def parse_table_file(data: bytes, filename: str | None) -> list[Row]:
    """CSV или XLSX — по расширению файла (или по сигнатуре zip у xlsx)."""
    if (filename or "").lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return parse_xlsx(data)
    return parse_csv(data)


def rows_to_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(["" if v is None else v for v in row] for row in rows)
    # BOM — чтобы Excel открыл кириллицу без танцев с кодировкой.
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------------- Google Sheets

def _service_account_info(raw: str) -> dict[str, Any]:
    raw = raw.strip()
    if not raw.startswith("{"):
        raw = base64.b64decode(raw).decode("utf-8")
    return json.loads(raw)


class SheetsClient:
    def __init__(self, settings: Settings) -> None:
        if not settings.sheets_enabled:
            raise SheetsError("Google Sheets не настроен: нужны SHEET_ID и GOOGLE_SERVICE_ACCOUNT_JSON")
        self.settings = settings
        self._spreadsheet = None
        self.last_read_title = ""

    def _open(self):
        if self._spreadsheet is None:
            import gspread

            client = gspread.service_account_from_dict(_service_account_info(self.settings.google_service_account_json))
            client.set_timeout(self.settings.sheets_timeout_seconds)
            self._spreadsheet = client.open_by_key(self.settings.sheet_id)
        return self._spreadsheet

    async def _call(self, fn: Callable[[], T], what: str) -> T:
        timeout = self.settings.sheets_timeout_seconds + 5
        try:
            async for attempt in AsyncRetrying(
                stop=stop_after_attempt(4),
                wait=wait_exponential(multiplier=2, min=2, max=30),
                retry=retry_if_not_exception_type(SheetMissing),
                reraise=True,
            ):
                with attempt:
                    if attempt.retry_state.attempt_number > 1:
                        log.warning("sheets_retry", what=what, attempt=attempt.retry_state.attempt_number)
                    return await asyncio.wait_for(asyncio.to_thread(fn), timeout)
        except SheetMissing:
            raise
        except Exception as exc:
            log.error("sheets_failed", what=what, error=repr(exc))
            raise SheetsError(f"Google Sheets: не удалось {what}: {exc}") from exc
        raise AssertionError("unreachable")

    SOURCE_NAMES = ("кандидаты", "candidates", "заявки", "applications")
    PASSED_ALIASES = ("прошли отбор", "прошедшие", "отобранные", "прошедшие отбор", "selected")

    def _worksheet(self, name: str):
        """Лист-источник заявок. Без имени: вкладка «Кандидаты» (или «Заявки»), иначе первая вкладка,
        которая не принадлежит боту («Статус бота», «Чекин …», «calls»)."""
        spreadsheet = self._open()
        sheets = spreadsheet.worksheets()
        if name:
            wanted = {name.strip().lower()}
            if name.strip().lower() in self.PASSED_ALIASES:
                wanted |= set(self.PASSED_ALIASES)
            for ws in sheets:
                if ws.title.strip().lower() in wanted:
                    return ws
            titles = ", ".join(f"«{ws.title}»" for ws in sheets)
            raise SheetMissing(f"В Google-таблице нет вкладки «{name}». Есть: {titles}. Создайте её и перенесите туда прошедших.")
        for ws in sheets:
            if ws.title.strip().lower() in self.SOURCE_NAMES:
                return ws
        own = {self.settings.sheet_status_worksheet.lower(), "calls"}
        for ws in sheets:
            title = ws.title.strip().lower()
            if title not in own and not title.startswith("чекин "):
                return ws
        return sheets[0]

    async def read_rows(self, worksheet: str) -> list[Row]:
        def fetch() -> list[Row]:
            ws = self._worksheet(worksheet)
            self.last_read_title = ws.title
            values = ws.get_all_values()
            if not values:
                return []
            header = [h.strip() for h in values[0]]
            rows = []
            for line in values[1:]:
                line = line + [""] * (len(header) - len(line))
                row = {h: v.strip() for h, v in zip(header, line) if h}
                if any(row.values()):
                    rows.append(row)
            return rows

        return await self._call(fetch, f"прочитать лист «{worksheet or 'первый'}»")

    async def write_tables(self, tables: list[Any]) -> None:
        """Перезаписывает несколько листов (объекты с полями name, header, rows) — по одному пакету на лист."""
        for table in tables:
            await self.write_table(table.name, table.header, table.rows)

    async def write_table(self, worksheet: str, header: list[str], rows: list[list[Any]]) -> None:
        """Полностью перезаписывает лист одним пакетом (создаёт лист, если его нет)."""
        values = [header, *[["" if v is None else v for v in row] for row in rows]]

        def write() -> None:
            import gspread

            spreadsheet = self._open()
            try:
                ws = spreadsheet.worksheet(worksheet)
            except gspread.WorksheetNotFound:
                ws = spreadsheet.add_worksheet(worksheet, rows=len(values) + 50, cols=len(header))
            # Размер листа ровно под данные: старые лишние строки исчезают, новым хватает места.
            ws.resize(rows=max(len(values), 2), cols=max(len(header), 1))
            ws.clear()
            ws.update(values, "A1", value_input_option="RAW")

        await self._call(write, f"записать лист «{worksheet}»")

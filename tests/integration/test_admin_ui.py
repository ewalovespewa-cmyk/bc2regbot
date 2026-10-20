"""Кнопочная админка, формат таблицы кандидатов, Excel, вкладки Google Sheet, сброс отметок."""
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from aiogram.methods import SendDocument
from sqlalchemy import func, select

from app.config import Settings
from app.db.models import Admin, Checkin, CheckinMethod, EventDay, OutboundEvent, Participant, Role
from app.db.session import create_sessionmaker
from app.services import checkin as checkin_service
from app.services.exports import all_tables, tables_to_xlsx
from app.services.importer import is_selected_value
from app.services.sheet_sync import import_rows, push_status_to_sheet
from app.services.sources import parse_csv, parse_xlsx
from tests.harness import Harness

SUPER, ADMIN_ID, VOL = 900, 901, 902

# Шапка как в таблице «Business Cup 2.0 — кандидаты».
CANDIDATES = (
    "Состояние,Имя,Почта,Метка,Participation type,Position in team,Team name,Phone number,Telegram,University,"
    "Year of study,Major,Interested in a career at Zimran,Вложения,Нужно внимание,Создано,ID заявки,LinkedIn / CV,"
    "Previous AIESEC project,Career interests\n"
    "Прошёл отбор,Dias Saiynov,dsaiynov01@gmail.com,,Solo participant,,,77475805502,@dsaiynov,NU,Graduate,CS,Yes,,,"
    "2026-10-19 19:19,id-1,,AIESEC Recruitment,Product\n"
    "заявка подана,Not Yet,notyet@gmail.com,,Solo participant,,,87001112233,@notyet,NU,1,CS,Yes,,,2026-10-19 19:20,id-2,,,\n"
    "прошел,Bizhan Ashykhatov,bizhanash@gmail.com,,Already have a team,Captain,Composer+,+77773375588,@bizhanchikk,"
    "AITU,3,SE,No,,,2026-10-19 21:09,id-3,,Business Cup,\n"
    "Прошёл отбор,Alikhan Altayev,ondyyvfx@gmail.com,,Already have a team,Team member,composer+,+77769973618,"
    "87029997339,QAIRU,1,AI,No,,,2026-10-19 21:11,id-4,,,\n"
    "Прошёл отбор,Dias Saiynov,dsaiynov01@gmail.com,,Solo participant,,,77475805502,@dsaiynov,NU,Graduate,CS,Yes,,,"
    "2026-10-20 10:00,id-5,,AIESEC Recruitment,Product\n"
)


def test_selected_values():
    values = Settings(_env_file=None).sheet_selected_values
    for ok in ("Прошёл отбор", "прошел", "ПРОШЛА", "accepted", "Selected", "yes", "Одобрен"):
        assert is_selected_value(ok, values), ok
    for no in ("заявка подана", "", "отказ", "не прошел"):
        assert not is_selected_value(no, values), no


async def test_import_candidates_sheet_format(session):
    report = await import_rows(session, Settings(_env_file=None), parse_csv(CANDIDATES.encode()))
    await session.commit()
    assert (report.created, report.not_selected, report.duplicates) == (3, 1, 1)
    people = {p.email: p for p in (await session.scalars(select(Participant))).all()}
    dias, alikhan = people["dsaiynov01@gmail.com"], people["ondyyvfx@gmail.com"]
    assert dias.application_id == "id-5"  # последняя заявка
    assert dias.phone == "+77475805502" and dias.tg_username_form == "dsaiynov"
    assert alikhan.tg_username_form is None  # в поле Telegram был номер
    assert alikhan.team_role == "Team member" and alikhan.team.name == people["bizhanash@gmail.com"].team.name


async def _seed_checkins(session):
    await import_rows(session, Settings(_env_file=None), parse_csv(CANDIDATES.encode()))
    tz = ZoneInfo("Asia/Almaty")
    now = datetime.now(UTC)
    today = EventDay(date=now.astimezone(tz).date(), label="День 1", checkin_opens_at=now - timedelta(hours=1),
                     checkin_closes_at=now + timedelta(hours=1), is_checkin_enabled=True)
    other = EventDay(date=(now + timedelta(days=1)).astimezone(tz).date(), label="День 2",
                     checkin_opens_at=now + timedelta(days=1), checkin_closes_at=now + timedelta(days=1, hours=2))
    session.add_all([today, other])
    await session.flush()
    people = (await session.scalars(select(Participant).order_by(Participant.id))).all()
    for p in people:
        p.telegram_user_id = 5000 + p.id
        session.add(Checkin(participant_id=p.id, event_day_id=today.id, method=CheckinMethod.SELF_QR))
    session.add(Checkin(participant_id=people[0].id, event_day_id=other.id, method=CheckinMethod.VOLUNTEER, checked_by=VOL))
    await session.commit()
    return today, other


async def test_tables_and_excel(session):
    await _seed_checkins(session)
    tables = await all_tables(session, "bc_bot", "Asia/Almaty")
    assert [t.name for t in tables] == ["Статус бота", "Чекин День 1", "Чекин День 2"]
    assert len(tables[0].rows) == 3 and len(tables[1].rows) == 3 and len(tables[2].rows) == 1
    assert tables[2].rows[0][6] == "волонтёр"

    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(tables_to_xlsx(tables)))
    assert wb.sheetnames == ["Статус бота", "Чекин День 1", "Чекин День 2"]
    assert parse_xlsx(tables_to_xlsx(tables))[0]["Ссылка-приглашение"].startswith("https://t.me/bc_bot?start=")


async def test_push_to_sheet_writes_all_tabs(session, monkeypatch):
    await _seed_checkins(session)
    written = []

    class FakeClient:
        def __init__(self, settings):
            pass

        async def write_tables(self, tables):
            written.extend(t.name for t in tables)

    monkeypatch.setattr("app.services.sheet_sync.SheetsClient", FakeClient)
    count = await push_status_to_sheet(session, Settings(_env_file=None, sheet_id="x", google_service_account_json="{}"))
    assert count == 3 and written == ["Статус бота", "Чекин День 1", "Чекин День 2"]


async def test_reset_and_delete_day(session):
    today, other = await _seed_checkins(session)
    assert await checkin_service.reset_checkins(session, today.id) == 3
    assert await checkin_service.checkins_count(session) == 1
    assert await session.scalar(select(func.count()).select_from(OutboundEvent).where(OutboundEvent.type == "checkin.cancelled")) == 3
    assert await checkin_service.delete_day(session, other) == 1
    assert await session.scalar(select(func.count()).select_from(EventDay)) == 1


# ---------------------------------------------------------------- кнопки


@pytest.fixture
async def h(engine, redis, session):
    session.add_all([Admin(telegram_user_id=ADMIN_ID, role=Role.ADMIN), Admin(telegram_user_id=VOL, role=Role.VOLUNTEER)])
    await session.commit()
    settings = Settings(_env_file=None, superadmin_ids=frozenset({SUPER}), checkin_secret="s", bot_username="bc_bot",
                        webhook_base_url="https://bot.example.com", screen_secret="k1", dashboard_secret="k2")
    return Harness(redis=redis, sessionmaker=create_sessionmaker(engine), settings=settings)


def _buttons(call) -> list[str]:
    return [b.callback_data or b.url for row in call.reply_markup.inline_keyboard for b in row]


def _last_with_markup(h, chat_id):
    return next(c for c in reversed(h.session.sent(chat_id)) if getattr(c, "reply_markup", None))


async def test_staff_start_opens_admin_panel(h):
    await h.text(SUPER, "/start")
    assert "m:checkin" in _buttons(_last_with_markup(h, SUPER))
    await h.text(VOL, "/start")
    assert _buttons(_last_with_markup(h, VOL)) == ["adm:find"]


async def test_import_by_file_buttons(h, session):
    async def download(file_id, destination, timeout):
        destination.write(CANDIDATES.encode())

    h.bot.download = download
    from aiogram.types import Chat, Document, Message, Update

    doc = Document(file_id="f1", file_unique_id="u1", file_name="Кандидаты.csv", file_size=len(CANDIDATES))
    msg = Message(message_id=1, date=datetime.now(UTC), chat=Chat(id=SUPER, type="private"), from_user=h.user(SUPER), document=doc)
    await h.feed(Update(update_id=80_001, message=msg))
    assert "a:file:selected" in _buttons(_last_with_markup(h, SUPER))
    await h.callback(SUPER, SUPER, "a:file:selected")
    assert "Новых участников: 3" in h.session.last_text(SUPER)
    assert await session.scalar(select(func.count()).select_from(Participant)) == 3


async def test_checkin_buttons_open_close_reset(h, session):
    today, _ = await _seed_checkins(session)
    today.is_checkin_enabled = False
    await session.commit()
    await h.callback(ADMIN_ID, ADMIN_ID, "m:checkin")
    assert "a:ci:open" in _buttons(_last_with_markup(h, ADMIN_ID))
    await h.callback(ADMIN_ID, ADMIN_ID, "a:ci:open")
    session.expire_all()
    assert await session.scalar(select(EventDay.is_checkin_enabled).where(EventDay.label == "День 1")) is True

    await h.callback(ADMIN_ID, ADMIN_ID, "a:resetok:all")  # админу сброс недоступен
    assert await checkin_service.checkins_count(session) == 4
    await h.callback(SUPER, SUPER, "m:reset")
    await h.callback(SUPER, SUPER, "a:reset:today")
    await h.callback(SUPER, SUPER, "a:resetok:today")
    assert await checkin_service.checkins_count(session) == 1


async def test_day_add_and_delete_buttons(h, session):
    await h.callback(SUPER, SUPER, "a:day:add")
    await h.text(SUPER, "19.10 00:00-23:59 Тест")
    day_id = await session.scalar(select(EventDay.id).where(EventDay.label == "Тест"))
    assert day_id
    await h.callback(SUPER, SUPER, f"a:day:delok:{day_id}")
    assert await session.scalar(select(func.count()).select_from(EventDay)) == 0


async def test_add_and_remove_admin_buttons(h, session):
    await h.text(4321, "/start", username="new_vol")  # бот запоминает username
    await h.callback(SUPER, SUPER, "a:adm:add")
    await h.text(SUPER, "@new_vol")
    await h.callback(SUPER, SUPER, "a:adm:role:volunteer")
    assert (await session.get(Admin, 4321)).role == Role.VOLUNTEER
    await h.callback(SUPER, SUPER, "a:adm:del:4321")
    session.expire_all()
    assert await session.get(Admin, 4321) is None


async def test_export_button_sends_excel(h, session):
    await _seed_checkins(session)
    await h.callback(ADMIN_ID, ADMIN_ID, "a:export")
    doc = next(c for c in h.session.calls if isinstance(c, SendDocument))
    assert "Статус бота" in doc.caption and "Чекин День 1" in doc.caption


async def test_poll_start_button(h, session):
    await h.callback(ADMIN_ID, ADMIN_ID, "a:poll:startok")
    from app.db.models import Poll

    assert await session.scalar(select(func.count()).select_from(Poll)) == 1
    assert "Опрос #1 запущен" in h.session.last_text(ADMIN_ID)


# ---------------------------------------------------------------- флоу «Кандидаты → Прошли отбор»

PASSED_HEADER = CANDIDATES.splitlines()[0]


def _passed_rows(*lines: str):
    # «Состояние» в скопированных строках может остаться «заявка подана» — во вкладке-источнике это не важно.
    return parse_csv(("\n".join([PASSED_HEADER, *lines]) + "\n").encode())


ROW_A = "заявка подана,Aru,aru@x.kz,,Solo participant,,,87011112233,@aru_b,NU,1,CS,Yes,,,2026-10-19 10:00,a1,,,"
ROW_B = "заявка подана,Dana,dana@x.kz,,Already have a team,Captain,Rockets,87022223344,@dana_o,NU,1,CS,Yes,,,2026-10-19 10:01,b1,,,"


async def test_list_mode_every_row_is_access_and_removed_rows_lose_it(session, monkeypatch):
    from app.services import sheet_sync

    sheet = {"rows": _passed_rows(ROW_A, ROW_B)}

    class FakeClient:
        last_read_title = "Прошли отбор"

        def __init__(self, settings):
            pass

        async def read_rows(self, worksheet):
            assert worksheet == "Прошли отбор"
            return sheet["rows"]

    monkeypatch.setattr(sheet_sync, "SheetsClient", FakeClient)
    settings = Settings(_env_file=None, sheet_id="x", google_service_account_json="{}")

    report, _ = await sheet_sync.import_from_sheet(session, settings)
    await session.commit()
    assert report.created == 2 and report.not_selected == 0  # «заявка подана» не мешает

    sheet["rows"] = _passed_rows(ROW_B)  # Aru убрали из вкладки
    report, _ = await sheet_sync.import_from_sheet(session, settings)
    await session.commit()
    assert report.deselected == 1
    session.expire_all()
    assert await session.scalar(select(Participant.is_selected).where(Participant.email == "aru@x.kz")) is False

    sheet["rows"] = _passed_rows(ROW_A, ROW_B)  # вернули — доступ снова есть, ссылка прежняя
    await sheet_sync.import_from_sheet(session, settings)
    await session.commit()
    session.expire_all()
    assert await session.scalar(select(Participant.is_selected).where(Participant.email == "aru@x.kz")) is True


async def test_list_mode_guard_against_wiped_tab(session, monkeypatch):
    from app.services import sheet_sync

    rows = [f"заявка подана,P{i},p{i}@x.kz,,Solo participant,,,8700000{i:04d},,NU,1,CS,Yes,,,,id{i},,," for i in range(30)]
    state = {"rows": _passed_rows(*rows)}

    class FakeClient:
        last_read_title = "Прошли отбор"

        def __init__(self, settings):
            pass

        async def read_rows(self, worksheet):
            return state["rows"]

    monkeypatch.setattr(sheet_sync, "SheetsClient", FakeClient)
    settings = Settings(_env_file=None, sheet_id="x", google_service_account_json="{}")
    await sheet_sync.import_from_sheet(session, settings)
    await session.commit()

    state["rows"] = _passed_rows(rows[0])  # случайно удалили почти всех
    report, _ = await sheet_sync.import_from_sheet(session, settings)
    await session.commit()
    assert report.deselected == 0 and report.deselect_blocked == 29
    assert await session.scalar(select(func.count()).select_from(Participant).where(Participant.is_selected)) == 30
